# -*- coding: utf-8 -*-
"""极简 Markdown → Word 转换器（够用版：标题/段落/图片/列表/引用/代码/表格/加粗）"""
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), ".pkgs"))

from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH


def set_font(style_or_run, name="微软雅黑", size=None):
    font = style_or_run.font
    font.name = name
    if size:
        font.size = Pt(size)
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run, "element") else None
    if rpr is not None:
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), name)


def add_inline(par, text):
    """支持 **加粗** 和 `代码`"""
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for p in parts:
        if not p:
            continue
        if p.startswith("**") and p.endswith("**"):
            r = par.add_run(p[2:-2])
            r.bold = True
        elif p.startswith("`") and p.endswith("`"):
            r = par.add_run(p[1:-1])
            r.font.name = "Consolas"
        else:
            par.add_run(p)


def convert(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    normal = doc.styles["Normal"]
    set_font(normal, "微软雅黑", 10.5)
    normal.paragraph_format.space_after = Pt(6)
    for i in range(1, 4):
        try:
            set_font(doc.styles["Heading %d" % i], "微软雅黑")
        except KeyError:
            pass

    base = os.path.dirname(os.path.abspath(md_path))
    i = 0
    in_code = False
    code_buf = []

    while i < len(lines):
        line = lines[i].rstrip()

        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                r = p.add_run("\n".join(code_buf))
                r.font.name = "Consolas"
                r.font.size = Pt(9)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2).strip(), level=level)
            i += 1
            continue

        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", line)
        if m:
            img = os.path.join(base, m.group(2))
            if os.path.exists(img):
                try:
                    doc.add_picture(img, width=Cm(15.5))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                except Exception:
                    doc.add_paragraph("[图片：%s]" % m.group(1))
            else:
                doc.add_paragraph("[图片缺失：%s]" % m.group(2))
            i += 1
            continue

        if line.lstrip().startswith("|") and "|" in line:
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            rows = [r for r in rows if not all(re.fullmatch(r":?-{2,}:?", c or "--") for c in r)]
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row[:len(rows[0])]):
                        table.cell(ri, ci).text = cell
            continue

        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            r = p.add_run(line.lstrip("> ").strip())
            r.italic = True
            i += 1
            continue

        m = re.match(r"^[-*]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(1))
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(2))
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, line)
        i += 1

    doc.save(docx_path)
    print("saved:", docx_path)


if __name__ == "__main__":
    for src in ["图文教程", "创作说明"]:
        md = os.path.join(r"D:\python\docs", src + ".md")
        if os.path.exists(md):
            convert(md, os.path.join(r"D:\python\docs", src + ".docx"))
