# -*- coding: utf-8 -*-
"""把 Markdown 排成 PPTX（16:9）。
# 一级标题 → 封面/章节；## 二级标题 → 分节页；- 列表 → 要点；普通段落 → 正文。
用法: python gen_ppt.py <输出pptx> <输入md1> [md2] ...
"""
import os, re, sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

DARK = RGBColor(0x1a, 0x2b, 0x4a)      # 深蓝，标题
ACCENT = RGBColor(0x2b, 0x6c, 0xb0)    # 亮蓝，章节
GRAY = RGBColor(0x66, 0x66, 0x66)      # 灰，副标题
MAX_BULLETS = 8                        # 一页最多几条要点


def _bullet(text):
    """去掉列表前缀，返回纯文本"""
    t = text.strip()
    t = re.sub(r"^([-*•]|\d+[.、])\s*", "", t)
    return t


def md_to_ppt(out, md_lines, title=None, author="花卷生成"):
    """把 markdown 行列表排成 PPTX，返回页数。"""
    # ---- 1. 解析成 (章节, [要点]) ----
    cover = title
    sections, cur = [], None
    for line in md_lines:
        st = line.strip()
        if not st:
            continue
        if st.startswith("# "):                      # 一级标题 = 封面/章节
            if cover is None:
                cover = st[2:].strip()
            if cur is not None and cur[1]:
                sections.append(cur)
            cur = (st[2:].strip(), [])
        elif st.startswith("## "):                   # 二级标题 = 新分节
            if cur is not None and cur[1]:
                sections.append(cur)
            cur = (st[2:].strip(), [])
        elif cur is not None:                        # 正文/要点
            cur[1].append(_bullet(st))
        else:
            cur = ("", [st])                         # 标题之前的内容先攒着
    if cur is not None:
        sections.append(cur)

    # ---- 2. 建 PPT ----
    prs = Presentation()
    prs.slide_width = Inches(13.333)                 # 16:9
    prs.slide_height = Inches(7.5)
    BLANK = prs.slide_layouts[6]

    # 封面
    s = prs.slides.add_slide(BLANK)
    tb = s.shapes.add_textbox(Inches(1), Inches(2.6), Inches(11.3), Inches(2.0))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = cover or "未命名"; p.alignment = PP_ALIGN.CENTER
    p.runs[0].font.size = Pt(40); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = DARK
    if author:
        p2 = tf.add_paragraph(); p2.text = author; p2.alignment = PP_ALIGN.CENTER
        p2.runs[0].font.size = Pt(18); p2.runs[0].font.color.rgb = GRAY

    # 章节页
    for (head, items) in sections:
        groups = [items[i:i + MAX_BULLETS] for i in range(0, len(items), MAX_BULLETS)] or [[]]
        for g in groups:
            s = prs.slides.add_slide(BLANK)
            hb = s.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.7), Inches(0.95))
            hf = hb.text_frame; hf.word_wrap = True
            hp = hf.paragraphs[0]; hp.text = head
            hp.runs[0].font.size = Pt(30); hp.runs[0].font.bold = True; hp.runs[0].font.color.rgb = ACCENT
            body = s.shapes.add_textbox(Inches(0.8), Inches(1.75), Inches(11.7), Inches(5.2))
            bf = body.text_frame; bf.word_wrap = True
            for i, it in enumerate(g):
                bp = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
                bp.text = ("• " + it) if it else ""
                bp.level = 0; bp.space_after = Pt(10)
                bp.runs[0].font.size = Pt(20)

    prs.save(out)
    return len(prs.slides._sldIdLst)


if __name__ == "__main__":
    out = sys.argv[1]
    md = []
    for p in sys.argv[2:]:
        with open(p, encoding="utf-8") as f:
            md.extend(f.read().splitlines())
    n = md_to_ppt(out, md)
    print("生成成功: %s  共 %d 页" % (out, n))
