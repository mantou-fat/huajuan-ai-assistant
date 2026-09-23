# -*- coding: utf-8 -*-
"""花卷文档生成 Skill 的自检：一条命令验证三件套真能生成、产物有效、能打开。
用法: python selfcheck.py
"""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "scripts"))
from gen_pdf import md_to_pdf
from gen_ppt import md_to_ppt
from gen_docx import md_to_docx

MD = ["# 自检文档", "",
      "## 一、要点", "- 要点一", "- 要点二", "",
      "## 二、清单", "| 列A | 列B |", "|---|---|", "| 1 | 2 |", "",
      "```", "print('hi')", "```"]


def main():
    out = ".selfcheck_tmp"
    os.makedirs(out, exist_ok=True)
    pdf = os.path.join(out, "t.pdf")
    pptx = os.path.join(out, "t.pptx")
    docx = os.path.join(out, "t.docx")
    results = []

    # 1) PDF：生成 + 文件头是 %PDF（不是空文件/坏文件）
    try:
        n = md_to_pdf(pdf, MD, title="自检")
        head = open(pdf, "rb").read(4)
        ok = n >= 1 and head == b"%PDF"
        results.append(("PDF 生成且是合法文件", ok, "共 %d 页" % n))
    except Exception as e:
        results.append(("PDF 生成且是合法文件", False, str(e)))

    # 2) PPT：生成 + 能被 python-pptx 打开、有页
    try:
        from pptx import Presentation
        md_to_ppt(pptx, MD, title="自检")
        slides = len(Presentation(pptx).slides)
        results.append(("PPT 生成且能打开", slides >= 1, "%d 页" % slides))
    except Exception as e:
        results.append(("PPT 生成且能打开", False, str(e)))

    # 3) Word：生成 + 能被 python-docx 打开、有段落
    try:
        from docx import Document
        md_to_docx(docx, MD, title="自检")
        paras = len(Document(docx).paragraphs)
        results.append(("Word 生成且能打开", paras >= 1, "%d 段" % paras))
    except Exception as e:
        results.append(("Word 生成且能打开", False, str(e)))

    # 清理临时产物
    for f in (pdf, pptx, docx):
        if os.path.exists(f):
            os.remove(f)
    try:
        os.rmdir(out)
    except OSError:
        pass

    passed = sum(1 for r in results if r[1])
    print("=" * 44)
    for name, ok, detail in results:
        print("[%s] %s  |  %s" % ("PASS" if ok else "FAIL", name, detail))
    print("=" * 44)
    print("自检结果: 通过 %d / 共 %d" % (passed, len(results)))
    print("✅ 三件套都正常，这包是验过的" if passed == len(results) else "❌ 有失败项，先修好再上架")


if __name__ == "__main__":
    main()
